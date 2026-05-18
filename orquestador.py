"""
orquestador.py

Pipeline principal del RPA.
Lee las filas pendientes del Cuadro Maestro y las pasa
por cada transaccion SAP en orden. Escribe el resultado
en la columna ESTADO del Excel al finalizar.
"""

import sys
import time
import importlib.util
import openpyxl
from datetime import datetime
from pathlib import Path
from dotenv import load_dotenv
from loguru import logger

load_dotenv()
sys.path.insert(0, str(Path(__file__).parent))

from Transacciones.Loguin   import ejecutar_login
from helpers.cuadro_maestro import leer_sin_estado_ok, RUTA_CUADRO_MAESTRO

# ---------------------------------------------------------------------------
# CONFIGURACION DE LOGS
# ---------------------------------------------------------------------------
LOG_DIR = Path(__file__).parent / "output" / "logs"
LOG_DIR.mkdir(parents=True, exist_ok=True)

logger.remove()
logger.add(
    sys.stdout,
    format="<green>{time:HH:mm:ss}</green> | <level>{level:<8}</level> | {message}",
    level="DEBUG",
)
logger.add(
    LOG_DIR / f"orquestador_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log",
    format="{time:YYYY-MM-DD HH:mm:ss} | {level:<8} | {message}",
    level="DEBUG",
    encoding="utf-8",
)


# ---------------------------------------------------------------------------
# PIPELINE: (nombre_legible, archivo, nombre_funcion)
# El orden aqui define el orden de ejecucion para cada fila.
# ---------------------------------------------------------------------------
PIPELINE = [
    ("ZSD007",     "T_ZSD007",     "ejecutar_ZSD007"),
    ("VA02",       "T_va02",       "ejecutar_VA02"),
    ("VL02N",      "T_vl02n",      "ejecutar_VL02N"),
    ("VF01",       "T_vf01",       "ejecutar_VF01"),
    ("ZEDOCPE001", "T_ZEDOCPE001", "ejecutar_ZEDOCPE001"),
]

COL_ESTADO  = "ESTADO"
COL_COD_PED = "COD PED"


# ---------------------------------------------------------------------------
# UTILIDADES
# ---------------------------------------------------------------------------

def _cargar_transaccion(archivo: str, funcion: str):
    """Carga dinamicamente una transaccion desde Transacciones/T_XXXX.py."""
    ruta = Path(__file__).parent / "Transacciones" / f"{archivo}.py"
    spec = importlib.util.spec_from_file_location(archivo, ruta)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return getattr(modulo, funcion)


def _escribir_estado(ruta: str, numero_fila_excel: int, valor: str) -> None:
    """Escribe el valor en la columna ESTADO de la fila indicada en el Excel."""
    wb = openpyxl.load_workbook(ruta)
    ws = wb["maestro"]

    col_idx = None
    for col in range(1, ws.max_column + 1):
        if str(ws.cell(5, col).value).strip() == COL_ESTADO:
            col_idx = col
            break

    if col_idx is None:
        logger.warning(f"Columna '{COL_ESTADO}' no encontrada en el Excel.")
        wb.close()
        return

    ws.cell(numero_fila_excel, col_idx).value = valor
    wb.save(ruta)
    wb.close()
    logger.debug(f"Excel actualizado → fila {numero_fila_excel} | {COL_ESTADO} = {valor}")


# ---------------------------------------------------------------------------
# PIPELINE POR FILA
# ---------------------------------------------------------------------------

def ejecutar_pipeline(session, fila: dict, numero_fila_excel: int) -> bool:
    """Ejecuta todas las transacciones del PIPELINE para una fila."""
    cod_ped = str(fila.get(COL_COD_PED, "?"))
    cliente = str(fila.get("CLIENTE", ""))

    logger.info("=" * 55)
    logger.info(f"FILA {numero_fila_excel} | Pedido: {cod_ped} | Cliente: {cliente}")
    logger.info("=" * 55)

    pasos_ok   = []
    pasos_fail = []

    for nombre, archivo, nombre_fn in PIPELINE:
        inicio = time.time()
        logger.info(f"[PASO] Iniciando {nombre}...")

        try:
            funcion   = _cargar_transaccion(archivo, nombre_fn)
            resultado = funcion(session, fila)
            duracion  = round(time.time() - inicio, 1)

            if resultado is False:
                logger.error(f"[PASO] {nombre} retornó False ({duracion}s) — deteniendo fila.")
                pasos_fail.append(nombre)
                _escribir_estado(RUTA_CUADRO_MAESTRO, numero_fila_excel, f"ERROR_{nombre}")
                _log_resumen_fila(cod_ped, pasos_ok, pasos_fail)
                return False

            logger.success(f"[PASO] {nombre} completado en {duracion}s.")
            pasos_ok.append(nombre)

        except Exception as e:
            duracion = round(time.time() - inicio, 1)
            logger.error(f"[PASO] {nombre} excepción ({duracion}s): {e}")
            pasos_fail.append(nombre)
            _escribir_estado(RUTA_CUADRO_MAESTRO, numero_fila_excel, f"ERROR_{nombre}")
            _log_resumen_fila(cod_ped, pasos_ok, pasos_fail)
            return False

        time.sleep(1)

    _escribir_estado(RUTA_CUADRO_MAESTRO, numero_fila_excel, "OK")
    _log_resumen_fila(cod_ped, pasos_ok, pasos_fail)
    return True


def _log_resumen_fila(cod_ped: str, ok: list, fail: list) -> None:
    logger.info(f"--- Resumen pedido {cod_ped} ---")
    for paso in ok:
        logger.info(f"   ✓ {paso}")
    for paso in fail:
        logger.error(f"   ✗ {paso}")


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def main():
    logger.info("Iniciando RPA Orquestador")
    logger.info(f"Cuadro Maestro: {RUTA_CUADRO_MAESTRO}")

    pendientes = leer_sin_estado_ok()
    logger.info(f"Filas pendientes encontradas: {len(pendientes)}")

    if pendientes.empty:
        logger.info("No hay filas pendientes. Proceso terminado.")
        return

    logger.info("Iniciando sesion SAP...")
    session = ejecutar_login()
    if not session:
        logger.error("No se pudo iniciar sesion SAP. Abortando.")
        return
    logger.success("Sesion SAP iniciada.")

    exitosas = 0
    errores  = 0

    for df_idx, fila in pendientes.iterrows():
        numero_fila_excel = df_idx + 6
        ok = ejecutar_pipeline(session, fila.to_dict(), numero_fila_excel)
        if ok:
            exitosas += 1
        else:
            errores += 1

    logger.info("=" * 55)
    logger.info("RESUMEN FINAL")
    logger.info(f"  Exitosas : {exitosas}")
    logger.info(f"  Errores  : {errores}")
    logger.info(f"  Total    : {exitosas + errores}")
    logger.info("=" * 55)
    logger.info(f"Log guardado en: {LOG_DIR}")


if __name__ == "__main__":
    main()
