import sys
import importlib
import importlib.util
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))

from Transacciones.Loguin import ejecutar_login

RUTA_EXCEL = Path(__file__).parent / "input" / "Facturas.xlsx"


def importar_transaccion(nombre_archivo, nombre_funcion):
    spec = importlib.util.spec_from_file_location(
        nombre_archivo,
        Path(__file__).parent / "Transacciones" / f"{nombre_archivo}.py"
    )
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return getattr(modulo, nombre_funcion)


def leer_ordenes():
    wb = openpyxl.load_workbook(RUTA_EXCEL)
    ws = wb.active

    encabezados = [str(ws.cell(1, c).value).strip().lower() for c in range(1, ws.max_column + 1)]

    col_orden  = encabezados.index("orden")  + 1
    col_estado = encabezados.index("estado") + 1

    filas = []
    for fila in range(2, ws.max_row + 1):
        orden = ws.cell(fila, col_orden).value
        if orden:
            filas.append({"fila": fila, "orden": orden})

    wb.close()
    return filas, col_estado


def marcar_estado(fila_num, col_estado, estado):
    wb = openpyxl.load_workbook(RUTA_EXCEL)
    ws = wb.active
    ws.cell(fila_num, col_estado).value = estado
    wb.save(RUTA_EXCEL)
    wb.close()


def ejecutar_prueba():
    print("=" * 50)
    print("  PRUEBA DE FLUJO COMPLETO")
    print("=" * 50)

    filas, col_estado = leer_ordenes()
    print(f"\nOrdenes encontradas en Excel: {len(filas)}")
    for f in filas:
        print(f"  Fila {f['fila']} -> Orden: {f['orden']}")

    session = ejecutar_login()
    if not session:
        print("ERROR: No se pudo iniciar sesion SAP")
        return
    
    print("\n--- Iniciando secuencia de transacciones ---\n")

    transacciones = [
        ("T-ZSD007",     "ejecutar_ZSD007"),
        ("T-va02",       "ejecutar_VA02"),
    ]

    resultados = {}
    todo_ok = True

    for archivo, funcion in transacciones:
        nombre = archivo.replace("T-", "")
        try:
            ejecutar = importar_transaccion(archivo, funcion)
            ok = ejecutar(session)
            resultados[nombre] = ok
            if not ok:
                todo_ok = False
        except Exception as e:
            print(f"[{nombre}] Error al importar o ejecutar: {e}")
            resultados[nombre] = False
            todo_ok = False

    print("\n" + "=" * 50)
    print("  RESUMEN")
    print("=" * 50)
    for transaccion, ok in resultados.items():
        estado = "OK" if ok else "FALLO"
        print(f"  {transaccion:<12} -> {estado}")

    if todo_ok:
        print("\nActualizando Excel...")
        for f in filas:
            marcar_estado(f["fila"], col_estado, "OK")
            print(f"  Fila {f['fila']} (Orden {f['orden']}) -> estado = OK")
        print("Excel actualizado correctamente.")
    else:
        print("\nHubo errores en el flujo, no se actualizo el Excel.")


if __name__ == "__main__":
    ejecutar_prueba()
